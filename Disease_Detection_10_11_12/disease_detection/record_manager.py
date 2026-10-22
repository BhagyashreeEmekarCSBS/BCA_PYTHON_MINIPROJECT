# record_manager.py
# Saves every patient's details and prediction to an Excel file.

import os
from datetime import datetime

from openpyxl import Workbook, load_workbook


class PatientRecordManager:

    HEADERS = [
        "Date and Time",
        "Patient Name",
        "Pregnancies",
        "Glucose",
        "Blood Pressure",
        "Skin Thickness",
        "Insulin",
        "BMI",
        "Diabetes Pedigree Function",
        "Age",
        "Prediction",
        "Diabetes Chance (%)",
    ]

    def __init__(self, file_name="patient_records.xlsx"):
        project_folder = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        self.file_path = os.path.join(project_folder, file_name)

    def save(self, name, values, result, probability):
        row = [
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            name,
            values[0],
            values[1],
            values[2],
            values[3],
            values[4],
            values[5],
            values[6],
            values[7],
            (
                "Likely to have Diabetes" if result == 1
                else "Not likely to have Diabetes"
            ),
            round(probability, 2),
        ]

        if os.path.exists(self.file_path):
            workbook = load_workbook(self.file_path)
            sheet = workbook.active
            sheet.append(row)
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Patient Records"
            sheet.append(self.HEADERS)
            sheet.append(row)

        workbook.save(self.file_path)
        return self.file_path
