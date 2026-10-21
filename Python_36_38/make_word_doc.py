import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_report():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_title.add_run("MINI PROJECT REPORT\n")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p_sub.add_run("AI-Driven Customer Segmentation for Retail Optimization")
    run2.font.size = Pt(14)
    run2.font.bold = True
    run2.font.color.rgb = RGBColor(68, 114, 196)

    doc.add_paragraph()

    p_team = doc.add_paragraph()
    p_team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = p_team.add_run(
        "Submitted by:\n"
        "1. Kartik More              PRN: RBC24BC036\n"
        "2. Shreya Nagargoje        PRN: RBC24BC038\n"
        "3. [Team Member 3]         PRN: [PRN No]\n"
        "4. [Team Member 4]         PRN: [PRN No]\n"
    )
    t.font.size = Pt(11)

    doc.add_paragraph("_" * 60)

    def section_header(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(31, 78, 121)

    def body(text):
        p = doc.add_paragraph(text)
        p.runs[0].font.size = Pt(11)

    # 1. Title
    section_header("1. Title of the Project")
    body("AI-Driven Customer Segmentation for Retail Optimization\n"
         "A machine learning system to divide retail shoppers into meaningful groups using K-Means Clustering, enabling smarter, targeted marketing strategies.")

    # 2. Problem Statement
    section_header("2. Problem Statement")
    body("In retail, every customer has different income and spending habits. Using a single generic discount strategy for all customers wastes marketing budget and reduces profit. "
         "The problem is: How can a retail store automatically identify different types of customers so it can give the right offer to the right customer at the right time?")

    # 3. Objective
    section_header("3. Objectives")
    for item in [
        "Load and clean the retail customer dataset (200 records).",
        "Select key features: Annual Income and Spending Score.",
        "Use the Elbow Method (WCSS) to find the best number of clusters (K = 5).",
        "Apply K-Means Clustering to group customers into 5 distinct segments.",
        "Provide clear, actionable retail marketing strategies for each segment."
    ]:
        doc.add_paragraph(item, style='List Bullet')

    # 4. Methodology
    section_header("4. Methodology")
    body(
        "Step 1 - Data Collection: Load Mall_Customers.csv containing 200 customer records.\n\n"
        "Step 2 - Feature Selection: Pick 'Annual Income (k$)' and 'Spending Score (1-100)' as the most relevant retail behavior indicators.\n\n"
        "Step 3 - Elbow Method: Compute WCSS (Within-Cluster Sum of Squares) for K = 1 to 10. "
        "The 'elbow' bend in the curve at K = 5 indicates the optimal cluster count.\n\n"
        "Step 4 - K-Means Clustering: Train KMeans(n_clusters=5) with k-means++ initialization. "
        "The algorithm assigns each customer to their nearest centroid and updates centroids iteratively until convergence.\n\n"
        "Step 5 - Visualization & Business Insights: Generate 2D cluster scatter plot and map each cluster to a retail marketing strategy."
    )

    # 5. Dataset, Tools & Technologies
    section_header("5. Dataset, Tools & Technologies Used")
    body(
        "Dataset: Mall Customers Dataset\n"
        "  - Records: 200 customers\n"
        "  - Attributes: CustomerID, Gender, Age, Annual Income (k$), Spending Score (1-100)\n\n"
        "Programming Language: Python 3.11\n\n"
        "Libraries Used:\n"
        "  - pandas     : Data loading and manipulation\n"
        "  - numpy      : Numerical operations\n"
        "  - scikit-learn: KMeans algorithm and Elbow Method\n"
        "  - matplotlib & seaborn: Data visualization and chart generation\n\n"
        "Tools: VS Code / Any Python IDE, Windows OS"
    )

    # 6. Implementation
    section_header("6. Implementation")
    body(
        "The project is implemented in a single clean Python file: customer_segmentation.py\n\n"
        "Key code steps:\n"
        "  A) Load dataset: df = pd.read_csv('Mall_Customers.csv')\n"
        "  B) Extract features: X = df.iloc[:, [3, 4]].values\n"
        "  C) Elbow loop: Fit KMeans for K=1 to 10, store inertia values in wcss[]\n"
        "  D) Train model: KMeans(n_clusters=5, init='k-means++', random_state=42).fit_predict(X)\n"
        "  E) Save outputs: elbow_method.png, customer_clusters.png, segmented_customers.csv"
    )

    # 7. Results
    section_header("7. Results")
    body("The K-Means model successfully identified 5 customer segments from 200 shoppers:")
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, h in enumerate(['Cluster Name', 'Count', 'Avg Income', 'Avg Spend', 'Retail Strategy']):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].font.bold = True

    for row_data in [
        ("Target / VIP", "39", "$86.5k", "82/100", "Premium loyalty, VIP perks"),
        ("Careful Savers", "35", "$88.2k", "17/100", "Bundle discounts, cashback"),
        ("Standard / Regular", "81", "$55.3k", "50/100", "Festival offers, membership"),
        ("Careless / Impulsive", "22", "$25.7k", "79/100", "Flash sales, EMI, BNPL"),
        ("Sensible / Budget", "23", "$26.3k", "21/100", "Clearance sales, essentials"),
    ]:
        r = table.add_row().cells
        for i, val in enumerate(row_data):
            r[i].text = val

    doc.add_paragraph()

    if os.path.exists("elbow_method.png"):
        p = doc.add_paragraph("Figure 1: Elbow Method Graph (Optimal K = 5)")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture("elbow_method.png", width=Inches(5.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    if os.path.exists("customer_clusters.png"):
        p = doc.add_paragraph("Figure 2: Customer Segmentation Cluster Plot")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture("customer_clusters.png", width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 8. Conclusion
    section_header("8. Conclusion")
    body(
        "This project successfully demonstrates how AI-driven K-Means Clustering can transform raw retail customer data into 5 meaningful business segments. "
        "Instead of using generic promotions for all customers, retail stores can now deliver personalized offers — "
        "VIP rewards for high spenders, budget deals for low-income shoppers, and trend-driven flash sales for impulsive buyers.\n\n"
        "This system reduces marketing costs, increases customer satisfaction, and boosts overall store revenue. "
        "It is a practical, scalable, and simple solution that showcases the power of Unsupervised Machine Learning in real-world retail business optimization."
    )

    doc.save("Mini_Project_Report.docx")
    print("Word report saved: Mini_Project_Report.docx")

create_report()
